from __future__ import annotations

import math
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
ASSET_DIR = ROOT / "report_assets"
REPORT_PATH = OUT_DIR / "NetSentry_AI_Internship_Report_Suleiman_Suleimanov.docx"

BLUE = "176B87"
DARK_BLUE = "12364A"
TEAL = "19A88B"
INK = "173042"
MUTED = "607987"
LIGHT = "F2F4F7"
PALE_TEAL = "EAF7F4"
WHITE = "FFFFFF"
RED = "B4233A"
AMBER = "A26400"


def font(size: int, bold: bool = False):
    path = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    return ImageFont.truetype(path, size)


def wrap_centered(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, face, fill: str, spacing: int = 8):
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 36
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=face)[2] <= max_width:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    heights = [draw.textbbox((0, 0), line, font=face)[3] for line in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=face)[2]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=face, fill=fill)
        y += height + spacing


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = "#5C8DA0", width: int = 5):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 14
    points = [
        end,
        (end[0] - size * math.cos(angle - .55), end[1] - size * math.sin(angle - .55)),
        (end[0] - size * math.cos(angle + .55), end[1] - size * math.sin(angle + .55)),
    ]
    draw.polygon(points, fill=color)


def create_architecture_image(path: Path) -> None:
    image = Image.new("RGB", (1600, 870), "#F7FAFC")
    draw = ImageDraw.Draw(image)
    draw.text((80, 55), "NetSentry AI — system architecture", font=font(42, True), fill="#173042")
    draw.text((82, 113), "Local, metadata-first monitoring pipeline", font=font(23), fill="#607987")

    boxes = {
        "sources": (80, 245, 355, 650),
        "processor": (485, 325, 760, 570),
        "analysis": (890, 205, 1195, 690),
        "interface": (1325, 325, 1535, 570),
    }
    for name, box in boxes.items():
        fill = {"sources": "#EAF7F4", "processor": "#EDF4FA", "analysis": "#F3EFFC", "interface": "#EAF7F4"}[name]
        outline = {"sources": "#19A88B", "processor": "#176B87", "analysis": "#7354A8", "interface": "#19A88B"}[name]
        draw.rounded_rectangle(box, radius=24, fill=fill, outline=outline, width=4)

    draw.text((122, 275), "INPUT SOURCES", font=font(23, True), fill="#177864")
    for i, item in enumerate(["Live interface", "PCAP import", "Synthetic demo"]):
        y = 350 + i * 82
        draw.rounded_rectangle((115, y, 320, y + 52), radius=12, fill="#FFFFFF", outline="#B8DCD4", width=2)
        wrap_centered(draw, (115, y, 320, y + 52), item, font(19), "#173042")

    wrap_centered(draw, boxes["processor"], "Packet processor\nMetadata parser", font(28, True), "#174E68", spacing=13)
    draw.text((928, 238), "ANALYSIS", font=font(23, True), fill="#63458F")
    for i, item in enumerate(["Thread-safe store", "Threat detector", "AI advisor"]):
        y = 320 + i * 105
        draw.rounded_rectangle((930, y, 1155, y + 68), radius=12, fill="#FFFFFF", outline="#CCBDE6", width=2)
        wrap_centered(draw, (930, y, 1155, y + 68), item, font(18), "#173042")

    wrap_centered(draw, boxes["interface"], "FastAPI and web dashboard", font(24, True), "#177864", spacing=11)
    arrow(draw, (355, 448), (485, 448))
    arrow(draw, (760, 448), (890, 448))
    arrow(draw, (1195, 448), (1325, 448))
    draw.text((500, 748), "Packet payloads are not retained; only selected metadata is stored in memory.", font=font(22), fill="#607987")
    image.save(path, quality=95)


def create_pipeline_image(path: Path) -> None:
    image = Image.new("RGB", (1600, 560), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((72, 50), "Detection and response workflow", font=font(40, True), fill="#173042")
    items = [
        ("1", "Capture", "Receive a packet or imported record"),
        ("2", "Normalise", "Extract addresses, ports, flags and protocol"),
        ("3", "Detect", "Apply sliding-window and state rules"),
        ("4", "Explain", "Create evidence and defensive guidance"),
        ("5", "Present", "Filter, inspect, ask AI and export"),
    ]
    x_positions = [70, 385, 700, 1015, 1330]
    for index, ((number, title, detail), x) in enumerate(zip(items, x_positions)):
        box = (x, 185, x + 235, 415)
        draw.rounded_rectangle(box, radius=20, fill="#F1F8F7" if index % 2 == 0 else "#F2F6FA", outline="#19A88B" if index % 2 == 0 else "#176B87", width=3)
        draw.ellipse((x + 86, 145, x + 149, 208), fill="#19A88B")
        num_width = draw.textbbox((0, 0), number, font=font(28, True))[2]
        draw.text((x + 117 - num_width / 2, 158), number, font=font(28, True), fill="#FFFFFF")
        wrap_centered(draw, (x + 12, 215, x + 223, 275), title, font(25, True), "#173042")
        wrap_centered(draw, (x + 15, 278, x + 220, 395), detail, font(16), "#607987", spacing=5)
        if index < len(items) - 1:
            arrow(draw, (x + 235, 300), (x_positions[index + 1], 300), width=4)
    image.save(path, quality=95)


def set_run(run, *, size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None, name: str = "Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def add_page_field(paragraph, total: bool = False):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " NUMPAGES " if total else " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instr, separate, text, end):
        run._r.append(item)
    set_run(run, size=9, color=MUTED)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblCellMar"))
    if existing is not None:
        tbl_pr.remove(existing)
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)


def set_table_geometry(table, widths: list[int], indent: int = 120):
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        node = tbl_pr.find(qn(tag))
        if node is not None:
            tbl_pr.remove(node)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
    set_cell_margins(table)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_custom_numbering(doc: Document, fmt: str, text: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    level.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
    if fmt == "bullet":
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Symbol")
        fonts.set(qn("w:hAnsi"), "Symbol")
        r_pr.append(fonts)
        level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


class ReportBuilder:
    def __init__(self) -> None:
        self.doc = Document()
        self.bullet_num = add_custom_numbering(self.doc, "bullet", "•")
        self.decimal_num = add_custom_numbering(self.doc, "decimal", "%1.")
        self._configure()

    def _configure(self):
        section = self.doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(.492)
        section.footer_distance = Inches(.492)
        section.different_first_page_header_footer = True

        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        normal.font.size = Pt(11)
        normal.font.color.rgb = RGBColor.from_string(INK)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.1
        normal.paragraph_format.widow_control = True

        heading_tokens = {
            "Heading 1": (16, BLUE, 16, 8),
            "Heading 2": (13, BLUE, 12, 6),
            "Heading 3": (12, DARK_BLUE, 8, 4),
        }
        for name, (size, color, before, after) in heading_tokens.items():
            style = styles[name]
            style.font.name = "Calibri"
            style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
            style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(color)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True

        caption = styles["Caption"]
        caption.font.name = "Calibri"
        caption.font.size = Pt(9)
        caption.font.italic = True
        caption.font.color.rgb = RGBColor.from_string(MUTED)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(4)
        caption.paragraph_format.space_after = Pt(10)

        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run("NETSENTRY AI  /  CYBERSECURITY INTERNSHIP PROJECT"), size=8.5, color=MUTED, bold=True)

        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run(p.add_run("Suleiman Suleimanov   •   Page "), size=9, color=MUTED)
        add_page_field(p)
        set_run(p.add_run(" of "), size=9, color=MUTED)
        add_page_field(p, total=True)

        settings = self.doc.settings.element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)

        core = self.doc.core_properties
        core.title = "NetSentry AI — Cybersecurity Internship Project Report"
        core.subject = "Network monitoring, anomaly detection and local AI assistance"
        core.author = "Suleiman Suleimanov"
        core.keywords = "cybersecurity, internship, packet analysis, IDS, FastAPI, Scapy, Ollama"
        core.comments = "Prepared as a project-based cybersecurity internship report."

    def page_break(self):
        self.doc.add_page_break()

    def heading(self, text: str, level: int = 1, new_page: bool = False):
        paragraph = self.doc.add_heading(text, level=level)
        if new_page:
            paragraph.paragraph_format.page_break_before = True
        return paragraph

    def para(self, text: str = "", *, bold_lead: str | None = None, italic: bool = False, align=None, after: float | None = None):
        p = self.doc.add_paragraph()
        if align is not None:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if bold_lead and text.startswith(bold_lead):
            set_run(p.add_run(bold_lead), bold=True)
            set_run(p.add_run(text[len(bold_lead):]), italic=italic)
        else:
            set_run(p.add_run(text), italic=italic)
        return p

    def bullet(self, text: str):
        p = self.doc.add_paragraph()
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), str(self.bullet_num))
        num_pr.extend([ilvl, num_id])
        set_run(p.add_run(text))
        return p

    def numbered(self, text: str, restart: bool = False):
        if restart:
            self.decimal_num = add_custom_numbering(self.doc, "decimal", "%1.")
        p = self.doc.add_paragraph()
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), str(self.decimal_num))
        num_pr.extend([ilvl, num_id])
        set_run(p.add_run(text))
        return p

    def callout(self, label: str, text: str, color: str = TEAL):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.2)
        p.paragraph_format.right_indent = Inches(.1)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(10)
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), PALE_TEAL)
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "20")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), color)
        borders.append(left)
        p_pr.extend([shd, borders])
        set_run(p.add_run(f"{label.upper()}  "), size=9.5, color=color, bold=True)
        set_run(p.add_run(text), size=10.5, color=INK)

    def code(self, text: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.18)
        p.paragraph_format.right_indent = Inches(.18)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(9)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "EDF1F4")
        p._p.get_or_add_pPr().append(shd)
        set_run(p.add_run(text), size=9, color=DARK_BLUE, name="Courier New")

    def table(self, headers: list[str], rows: list[list[str]], widths: list[int], font_size: float = 9.2):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        header_row = table.rows[0]
        repeat_header(header_row)
        for index, value in enumerate(headers):
            cell = header_row.cells[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            shade_cell(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run(p.add_run(value), size=font_size, color=DARK_BLUE, bold=True)
        for values in rows:
            row = table.add_row()
            for index, value in enumerate(values):
                cell = row.cells[index]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                set_run(p.add_run(str(value)), size=font_size, color=INK)
        set_table_geometry(table, widths)
        return table

    def figure(self, path: Path, caption: str, width: float = 6.25):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        shape = p.add_run().add_picture(str(path), width=Inches(width))
        shape._inline.docPr.set("descr", caption)
        self.doc.add_paragraph(caption, style="Caption")

    def cover(self, logo_path: Path):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = p.add_run().add_picture(str(logo_path), width=Inches(1.1))
        shape._inline.docPr.set("descr", "NetSentry AI shield and network-node mark")
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(12)
        set_run(p.add_run("CYBERSECURITY INTERNSHIP PROJECT REPORT"), size=10.5, color=TEAL, bold=True)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(7)
        set_run(p.add_run("NetSentry AI"), size=32, color=DARK_BLUE, bold=True)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(42)
        set_run(p.add_run("Network Traffic Monitoring, Suspicious Activity Detection,\nand Local AI-Assisted Security Guidance"), size=15, color=BLUE)
        self.para("Prepared by", align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        set_run(p.add_run("Suleiman Suleimanov"), size=16, color=INK, bold=True)
        self.para("Project-based cybersecurity internship • Eight-week implementation cycle", align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
        self.para("July 2026", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(54)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run("A local-first educational network monitoring system"), size=9.5, color=MUTED, italic=True)

    def add_toc(self, entries: list[tuple[str, int, int]] | None = None):
        self.heading("Table of Contents", 1)
        for title, page, level in entries or []:
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.22 * max(0, level - 1))
            p.paragraph_format.space_after = Pt(4 if level == 1 else 2)
            tabs = p.paragraph_format.tab_stops
            tabs.add_tab_stop(Inches(6.35), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            run = p.add_run(f"{title}\t{page}")
            set_run(run, size=10.5 if level == 1 else 9.5, color=INK if level == 1 else MUTED, bold=level == 1)


def create_logo(path: Path) -> None:
    image = Image.new("RGBA", (520, 520), (255, 255, 255, 0))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((70, 70, 450, 450), radius=92, fill="#EAF7F4", outline="#19A88B", width=10)
    points = [(260, 125), (375, 172), (356, 330), (260, 400), (164, 330), (145, 172)]
    draw.polygon(points, fill="#FFFFFF", outline="#176B87", width=10)
    draw.line((205, 250, 315, 250), fill="#19A88B", width=12)
    draw.line((260, 195, 260, 305), fill="#19A88B", width=12)
    for x, y in [(205, 250), (315, 250), (260, 195), (260, 305), (260, 250)]:
        draw.ellipse((x - 16, y - 16, x + 16, y + 16), fill="#176B87")
    image.save(path)


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    logo = ASSET_DIR / "netsentry_report_logo.png"
    architecture = ASSET_DIR / "netsentry_architecture.png"
    pipeline = ASSET_DIR / "netsentry_detection_pipeline.png"
    create_logo(logo)
    create_architecture_image(architecture)
    create_pipeline_image(pipeline)

    r = ReportBuilder()
    r.cover(logo)
    r.page_break()

    r.heading("Declaration", 1)
    r.para(
        "I declare that this report presents the design, implementation and evaluation of NetSentry AI, an educational network-monitoring application developed as my cybersecurity internship project. The technical descriptions, test results and reflections in this document correspond to the submitted project files. External frameworks and references are identified in the references section."
    )
    r.para("Name: Suleiman Suleimanov", bold_lead="Name:")
    r.para("Project: NetSentry AI", bold_lead="Project:")
    r.para("Date: July 2026", bold_lead="Date:")
    r.para("Signature: ______________________________", bold_lead="Signature:")

    r.heading("Acknowledgements", 1)
    r.para(
        "I would like to acknowledge the guidance, feedback and learning opportunities provided throughout the internship period. The project benefited from discussions about network protocols, secure software design, defensive monitoring and responsible handling of traffic data. I also appreciate the open-source communities behind Python, Scapy, FastAPI and Ollama, whose documentation made it possible to turn theoretical concepts into a working prototype."
    )

    r.heading("Executive Summary", 1)
    r.para(
        "NetSentry AI is a local network traffic monitoring and suspicious-activity detection application created during an eight-week project-based cybersecurity internship. The system addresses a practical learning problem: packet-analysis tools provide detailed evidence, but beginners often find it difficult to convert packet metadata into an investigation plan. NetSentry combines live capture, PCAP analysis, protocol filtering, explainable detection rules and an optional local AI advisor in one browser-based dashboard."
    )
    r.para(
        "The backend uses Python, Scapy and FastAPI. Traffic is normalised into a privacy-conscious metadata model containing timestamps, addresses, ports, protocols, lengths, selected TCP flags and DNS query names. Payloads are not retained. A thread-safe in-memory store supplies summary statistics and recent records to the interface. The detector evaluates sliding time windows and state changes to identify possible TCP port scans, SYN floods, ICMP floods, DNS-tunnelling indicators, ARP address changes and connections to risky service ports. Every alert includes evidence, severity and a defensive recommendation."
    )
    r.para(
        "The advisor can call a local Ollama model and automatically falls back to deterministic rule-based guidance when a model is unavailable. A synthetic demo mode provides repeatable results without administrator privileges or access to real traffic. Eight automated tests passed, JavaScript syntax validation succeeded, and an end-to-end local API smoke test processed 140 synthetic packets, 37,125 bytes and three expected alerts. The project demonstrates full-stack development, network analysis, anomaly detection, privacy-aware design, testing and technical documentation."
    )
    r.page_break()
    r.add_toc(
        [
            ("Declaration", 2, 1),
            ("Acknowledgements", 2, 1),
            ("Executive Summary", 2, 1),
            ("List of Abbreviations", 4, 1),
            ("1. Introduction", 5, 1),
            ("2. Internship Objectives and Requirements", 6, 1),
            ("3. Background and Technology Selection", 7, 1),
            ("4. System Design and Architecture", 8, 1),
            ("5. Implementation", 10, 1),
            ("6. Suspicious-Activity Detection", 11, 1),
            ("7. Local AI Advisor", 12, 1),
            ("8. Dashboard and User Workflow", 13, 1),
            ("9. Testing and Evaluation", 14, 1),
            ("10. Project Management and Weekly Work Record", 15, 1),
            ("11. Challenges and Solutions", 16, 1),
            ("12. Skills and Learning Outcomes", 17, 1),
            ("13. Ethics, Limitations and Future Work", 18, 1),
            ("14. Conclusion", 19, 1),
            ("References", 20, 1),
            ("Appendix A — Local Installation and Demonstration", 21, 1),
            ("Appendix B — Project File Structure", 22, 1),
        ]
    )
    r.page_break()

    r.heading("List of Abbreviations", 1)
    r.table(
        ["Abbreviation", "Meaning"],
        [
            ["AI", "Artificial Intelligence"], ["API", "Application Programming Interface"],
            ["ARP", "Address Resolution Protocol"], ["CSV", "Comma-Separated Values"],
            ["DNS", "Domain Name System"], ["HTTP", "Hypertext Transfer Protocol"],
            ["ICMP", "Internet Control Message Protocol"], ["IDS", "Intrusion Detection System"],
            ["IP", "Internet Protocol"], ["LLM", "Large Language Model"],
            ["MVP", "Minimum Viable Product"], ["PCAP", "Packet Capture file format"],
            ["TCP", "Transmission Control Protocol"], ["TLS", "Transport Layer Security"],
            ["UDP", "User Datagram Protocol"], ["UI", "User Interface"],
        ],
        [2100, 7260],
    )

    r.heading("1. Introduction", 1, new_page=True)
    r.heading("1.1 Project context", 2)
    r.para(
        "Network visibility is a basic requirement for defensive security. Administrators need to know which systems are communicating, which protocols are active, whether traffic volume has changed and whether repeated patterns could indicate reconnaissance, abuse or misconfiguration. Wireshark is an established packet analyser, but its breadth can be difficult for a learner who is still connecting protocol details to security decisions. This internship project explored how a smaller application could present the most useful metadata and explain why a pattern deserves attention."
    )
    r.para(
        "The result is NetSentry AI: a local application that captures or imports traffic, classifies common protocols, displays packet metadata, applies transparent anomaly rules and provides investigation advice. The name reflects the two main ideas of the system: continuous network observation and a defensive assistant that helps the user interpret findings."
    )
    r.heading("1.2 Problem statement", 2)
    r.para(
        "Packet data is highly detailed but does not automatically answer operational questions. A list of TCP SYN packets may be normal connection activity or may represent a scan. A long DNS query may be generated by a legitimate service or may indicate data encoded in a hostname. An ARP update may be routine or may indicate address spoofing. The project therefore needed to preserve evidence while avoiding unsupported claims."
    )
    r.callout("Detection principle", "NetSentry reports suspicious indicators, not confirmed attacks. Every result must be verified against authorised asset information and additional logs.")
    r.heading("1.3 Scope", 2)
    r.bullet("Capture packet metadata from a user-selected Linux interface through Scapy/libpcap.")
    r.bullet("Import authorised PCAP, PCAPNG or CAP files for offline analysis.")
    r.bullet("Identify common network and application protocols and provide interactive filters.")
    r.bullet("Detect a defined set of suspicious patterns with explainable thresholds.")
    r.bullet("Provide local AI-assisted or deterministic defensive guidance.")
    r.bullet("Operate as an educational MVP rather than a production replacement for Zeek, Suricata or an enterprise SIEM.")

    r.heading("2. Internship Objectives and Requirements", 1, new_page=True)
    r.heading("2.1 Learning objectives", 2)
    r.para(
        "The internship was organised around a complete development cycle rather than a collection of disconnected exercises. I intended to strengthen both cybersecurity knowledge and software engineering practice by moving from a problem statement to an evaluated deliverable."
    )
    for item in [
        "Understand how Ethernet, IP, TCP, UDP, DNS, ICMP and ARP information appears in captured traffic.",
        "Translate suspicious network behaviours into measurable, testable rules.",
        "Build a structured Python backend that remains stable while packets arrive from a separate capture thread.",
        "Create an interface that supports filtering and investigation instead of only displaying raw records.",
        "Integrate an LLM in a way that does not make core detection dependent on AI availability.",
        "Test the implementation with synthetic traffic and document limitations honestly.",
    ]:
        r.bullet(item)

    r.heading("2.2 Functional and non-functional requirements", 2)
    r.table(
        ["ID", "Requirement", "Acceptance criterion"],
        [
            ["FR-01", "Capture live traffic metadata", "User can choose live mode and an available interface."],
            ["FR-02", "Support repeatable demonstration", "Demo mode works without root and creates expected alerts."],
            ["FR-03", "Classify protocols", "Dashboard separates TCP, UDP, DNS, HTTP, TLS, SSH, FTP, Telnet, ICMP, ARP and IPv6."],
            ["FR-04", "Detect suspicious behaviour", "Rules produce an alert with severity, evidence and a recommendation."],
            ["FR-05", "Explore traffic", "User can search, filter, inspect a packet and show only suspicious records."],
            ["FR-06", "Analyse stored captures", "Authorised PCAP-family files can be imported with bounded size and record count."],
            ["FR-07", "Provide advice", "Local Ollama is used when available; rules answer when it is not."],
            ["NFR-01", "Privacy", "Payloads are not retained and the server binds to localhost by default."],
            ["NFR-02", "Resilience", "Malformed packets and unavailable AI do not stop monitoring."],
            ["NFR-03", "Usability", "Primary capture and investigation actions are visible from one dashboard."],
            ["NFR-04", "Testability", "Core rules and processing pipeline can be tested without network privileges."],
        ],
        [850, 3200, 5310],
        font_size=8.6,
    )

    r.heading("3. Background and Technology Selection", 1, new_page=True)
    r.heading("3.1 Network monitoring concepts", 2)
    r.para(
        "Passive monitoring observes network communication without initiating connections to the systems being examined. Packet metadata can reveal endpoints, service ports, timing, protocol distribution and connection flags. This information supports troubleshooting and intrusion detection, but it can also contain sensitive operational details. NetSentry therefore follows a metadata-minimisation approach and does not retain application payloads."
    )
    r.para(
        "Signature detection looks for known values or patterns, while anomaly detection compares behaviour with expected limits or baselines. The MVP uses explainable threshold and state rules. This choice supports an internship demonstration because each result can be traced to specific evidence. It also exposes an important limitation: one fixed threshold cannot represent every network."
    )
    r.heading("3.2 Selected technologies", 2)
    r.table(
        ["Technology", "Role", "Reason for selection"],
        [
            ["Python 3.10+", "Application language", "Readable syntax, strong networking ecosystem and rapid prototyping."],
            ["Scapy", "Capture and packet parsing", "Supports sniffing, protocol layers and PCAP reading through a Python API."],
            ["FastAPI", "Local web API", "Typed request models, automatic API documentation and straightforward endpoint design."],
            ["Pydantic", "Data validation", "Defines consistent packet, alert and request structures."],
            ["HTML/CSS/JavaScript", "Dashboard", "Runs locally without a separate front-end build pipeline."],
            ["Ollama", "Optional local LLM", "Keeps the explanation workflow local and supports a simple HTTP API."],
            ["unittest", "Automated testing", "Part of Python's standard library and suitable for privilege-free rule tests."],
        ],
        [1700, 2200, 5460],
        font_size=8.8,
    )
    r.heading("3.3 Technology trade-offs", 2)
    r.para(
        "A desktop GUI toolkit was considered, but a local web interface offered better responsive layout, simpler tables and easier separation between capture logic and presentation. A JavaScript chart dependency was not required; the packet-rate graph is drawn with the browser Canvas API, and protocol composition uses a CSS conic gradient. SQLite persistence was deferred because an in-memory store reduces the risk of retaining sensitive metadata during a classroom demonstration."
    )

    r.heading("4. System Design and Architecture", 1, new_page=True)
    r.figure(architecture, "Figure 1. NetSentry AI component architecture and local data flow.")
    r.heading("4.1 Component responsibilities", 2)
    r.para(
        "The capture controller owns one active source. In demo mode it runs a background generator; in live mode it wraps Scapy's asynchronous sniffer; during PCAP import it reads stored packets with a bounded limit. Every source sends records to the packet processor, which is the common entry point for storage and detection. This prevents different input modes from producing inconsistent results."
    )
    r.para(
        "The parser converts Scapy objects into a PacketRecord. The store assigns monotonically increasing identifiers, updates counters and retains recent packets and alerts in bounded deques. A re-entrant lock protects the structures because live capture and HTTP requests run in different execution contexts. The detector maintains its own sliding windows and cooldown state. FastAPI exposes sanitised snapshots to the browser, and the advisor receives only summary data and alerts."
    )
    r.heading("4.2 Data model", 2, new_page=True)
    r.table(
        ["Entity", "Important fields", "Purpose"],
        [
            ["PacketRecord", "id, time, src/dst, ports, protocol, length, direction, flags, DNS query", "Normalised metadata shared by capture, detection and UI."],
            ["AlertRecord", "type, title, severity, source, packet id, evidence, recommendation", "Explainable result linked to the triggering packet."],
            ["Capture status", "running, mode, interface, start time, last error", "Keeps UI controls consistent with backend state."],
            ["Snapshot", "statistics, filtered packets, recent alerts, capture status", "Single polling response used to refresh the dashboard."],
        ],
        [1700, 4560, 3100],
        font_size=8.8,
    )
    r.heading("4.3 Privacy and security design", 2)
    r.bullet("The default server address is 127.0.0.1, so the dashboard is not exposed to the network unless the operator changes the host option.")
    r.bullet("Packet payload bytes are neither added to PacketRecord nor exported to CSV.")
    r.bullet("The store is bounded to 5,000 recent packets and 1,000 recent alerts.")
    r.bullet("PCAP uploads are limited to 100 MB and 10,000 processed packets and are deleted after import.")
    r.bullet("The advisor prompt explicitly frames results as indicators and requests defensive rather than offensive guidance.")

    r.heading("5. Implementation", 1, new_page=True)
    r.heading("5.1 Packet acquisition", 2)
    r.para(
        "Live capture uses Scapy AsyncSniffer with store=False. Disabling Scapy's internal packet list avoids keeping a second copy of the traffic. The callback sends each packet to the processor and catches parsing errors so that one malformed or unsupported frame cannot terminate the capture. Interface names are obtained through the operating system socket API and presented in the dashboard."
    )
    r.para(
        "Linux restricts raw packet capture. The README therefore documents a clear distinction: demo and PCAP analysis can run as a normal user, while live capture normally requires administrator privileges or carefully assigned capture capabilities. This is an operating-system security control rather than an application defect."
    )
    r.heading("5.2 Packet parsing and protocol classification", 2)
    r.para(
        "The parser checks packet layers in a deliberate order. ARP requires link-layer addresses; TCP and UDP expose ports; ICMP exposes type and code; DNS queries override the generic UDP classification. Common service ports map to readable labels such as HTTP, TLS, SSH, FTP and Telnet. The parser also derives a simple direction label by comparing private and public IP-address status."
    )
    r.heading("5.3 Storage and concurrency", 2)
    r.para(
        "TrafficStore contains bounded deques, cumulative counters and a protocol Counter. Every mutation is performed under an RLock. Snapshot generation copies the current lists while locked and applies user filters after releasing the lock. This reduces the time the capture callback must wait while the browser searches records. Clearing data also resets identifiers and protocol counters."
    )
    r.heading("5.4 API and browser interface", 2)
    r.para(
        "FastAPI routes manage the capture lifecycle, return interfaces and snapshots, clear state, import PCAP files, export CSV and answer advisor requests. Pydantic constrains capture modes and question length. The browser polls one snapshot endpoint each second; this is simpler than WebSockets for the expected educational traffic volume and makes reconnection automatic."
    )
    r.table(
        ["Method and endpoint", "Function", "Important control"],
        [
            ["GET /api/snapshot", "Return statistics, filtered packets and alerts", "Limit 1–1,000 and query length limit"],
            ["POST /api/capture/start", "Start demo or live capture", "Rejects a second active capture"],
            ["POST /api/capture/stop", "Stop the active source", "Safe if already stopped"],
            ["POST /api/import/pcap", "Analyse a stored capture", "Extension, size and record limits"],
            ["POST /api/advisor", "Explain current context or one alert", "Question validation and defensive prompt"],
            ["GET /api/export.csv", "Download current metadata", "No payload field is available"],
            ["DELETE /api/data", "Clear packets, alerts and detector state", "Explicit user action"],
        ],
        [2550, 3450, 3360],
        font_size=8.6,
    )

    r.heading("6. Suspicious-Activity Detection", 1, new_page=True)
    r.figure(pipeline, "Figure 2. From captured metadata to an explainable defensive response.")
    r.heading("6.1 Detection strategy", 2)
    r.para(
        "The detector is intentionally transparent. It maintains short deques keyed by source or destination, removes records older than the current window and evaluates counts or uniqueness. A cooldown key prevents the same continuing behaviour from filling the alert list every second. This approach is efficient for the MVP and is straightforward to unit test."
    )
    r.table(
        ["Indicator", "Evidence and default threshold", "Window", "Severity", "Primary response"],
        [
            ["TCP port scan", "≥18 initial SYN packets and ≥10 unique destination/port pairs", "10 s", "High", "Verify source; review firewall follow-up; restrict if unknown."],
            ["SYN flood", "≥35 initial SYN packets directed to one destination", "5 s", "Critical", "Check availability; enable SYN protection/rate limits."],
            ["ICMP flood", "≥40 ICMP packets from one source", "5 s", "High", "Validate monitoring source; apply ICMP rate limits."],
            ["DNS tunnelling indicator", "Query ≥70 characters or label ≥40, with entropy ≥3.5", "Per query", "Medium", "Inspect repetition/domain; enforce approved resolvers."],
            ["ARP address change", "Known IP is advertised from a different MAC", "Stateful", "Critical", "Validate MAC; isolate conflict; use switch controls."],
            ["Risky service port", "Destination 23, 2323, 4444, 5555 or 31337", "Per flow", "Low", "Confirm authorisation; replace clear-text services."],
        ],
        [1600, 2910, 720, 850, 3280],
        font_size=7.9,
    )
    r.heading("6.2 False positives and interpretation", 2)
    paragraph = r.para(
        "A vulnerability scanner can resemble hostile reconnaissance; load testing can resemble a SYN flood; monitoring tools can produce frequent ICMP; content-delivery systems can use long DNS names; and legitimate network reconfiguration can change an ARP association. For this reason, alert wording uses possible and suspicious rather than confirmed. The recommended workflow is to establish whether the source is authorised, compare the timestamp with other logs and then decide whether containment is justified. All thresholds are demonstration defaults; a production version should use configurable, network-specific baselines and cooldowns."
    )
    paragraph.paragraph_format.keep_together = True

    r.heading("7. Local AI Advisor", 1)
    r.heading("7.1 Purpose and boundaries", 2)
    r.para(
        "The AI component is an explanation layer, not a detection authority. The detector creates the alert before any model is called. The model receives summary counters, at most five recent alerts and an optional selected alert. It does not receive packet payloads. The system prompt instructs it to preserve uncertainty, provide ordered defensive actions and answer in the user's language."
    )
    r.heading("7.2 Ollama integration", 2)
    r.para(
        "When enabled, the advisor sends a non-streaming chat request to the local Ollama endpoint, which defaults to http://127.0.0.1:11434. The model name is configurable through NETSENTRY_OLLAMA_MODEL and defaults to qwen2.5:3b. A short timeout prevents an unavailable model from blocking the application indefinitely."
    )
    r.heading("7.3 Deterministic fallback", 2)
    r.para(
        "If Ollama is missing, stopped or returns an error, the advisor generates a response from the selected alert's title, description, severity and recommendation. It adds standard verification and evidence-preservation steps. General questions receive a capture summary and prioritisation advice. This fallback made the demonstration reliable and also created a clear separation between deterministic security logic and probabilistic language generation."
    )
    r.callout("Safety boundary", "No alert depends on an LLM response. AI output is advice for a human investigator and is explicitly labelled as unable to confirm an attack.")
    r.heading("7.4 Prompt-injection consideration", 2)
    r.para(
        "Because payloads are not included, application data captured from the network cannot directly become instructions to the model. DNS query names and alert text remain untrusted contextual data, so the system prompt tells the model to use only supplied metadata and maintain a defensive role. A future version should serialise context more formally, escape untrusted values and apply output policy checks."
    )

    r.heading("8. Dashboard and User Workflow", 1, new_page=True)
    r.heading("8.1 Overview", 2)
    r.para(
        "The interface uses a dark security-operations visual style and is responsive for laptop and mobile widths. Four cards display total packets, traffic volume, alert count and engine status. A Canvas graph shows the most recent packet rate, while a protocol donut summarises the cumulative distribution. These elements provide context before the user opens individual records."
    )
    r.heading("8.2 Packet investigation", 2)
    r.numbered("Choose Synthetic demo or Live interface and start capture.", restart=True)
    r.numbered("Review packet rate, volume and protocol distribution for unexpected changes.")
    r.numbered("Filter by protocol or search for an address, port or detail string.")
    r.numbered("Enable Alerts only to isolate records associated with a detector result.")
    r.numbered("Open a row to inspect timestamp, direction, flags and DNS metadata.")
    r.numbered("Review the linked alert, evidence and severity before taking action.")
    r.numbered("Ask the advisor for a prioritised investigation plan and correlate with external logs.")
    r.heading("8.3 Demonstration mode", 2)
    r.para(
        "The generator produces ordinary TLS, DNS, TCP, UDP, ICMP and HTTP metadata. At deterministic phases it adds SYN sequences across multiple ports, a high-entropy DNS name and a connection to port 4444. This allows an evaluator to see highlighted packets and alerts without exposing real organisational traffic. The same PacketProcessor handles demo and real inputs, so the detection path remains representative."
    )
    r.heading("8.4 Data export", 2)
    r.para(
        "CSV export includes packet identifier, time, endpoints, ports, protocol, transport, length, direction, suspicious flag, severity and summary text. It does not contain payload bytes. The export supports later analysis in a spreadsheet or inclusion of selected results in a report, but it must still be handled as potentially sensitive operational metadata."
    )

    r.heading("9. Testing and Evaluation", 1, new_page=True)
    r.heading("9.1 Test strategy", 2)
    r.para(
        "Testing focused on the logic that could create incorrect security findings or break the local demonstration. Unit tests constructed PacketRecord objects directly so they did not require a network interface. Pipeline tests exercised identifier assignment, counters, protocol filters, suspicious filters and reset behaviour. A full local smoke test then started the FastAPI application in automatic demo mode and queried the health, snapshot, advisor and HTML endpoints."
    )
    r.table(
        ["Test area", "Cases", "Result", "Evidence"],
        [
            ["Threat detector", "Port scan; normal SYN; long/high-entropy DNS; ARP change; risky-port cooldown", "5/5 passed", "python -m unittest discover -s tests -v"],
            ["Processing pipeline", "Identifiers/counters; filters; clear/reset", "3/3 passed", "Same automated test command"],
            ["Python syntax", "Application and tests", "Passed", "python -m compileall -q netsentry tests"],
            ["JavaScript syntax", "Dashboard controller", "Passed", "node --check netsentry/static/app.js"],
            ["API health", "GET /api/health", "HTTP 200", "Returned status=ok and version=0.1.0"],
            ["Demo snapshot", "Eight-second controlled run", "Passed", "140 packets; 37,125 bytes; 3 alerts"],
            ["Advisor fallback", "Ollama intentionally disabled", "Passed", "Returned rule-based prioritisation"],
            ["Application page", "GET /", "HTTP 200", "HTML title returned NetSentry AI"],
        ],
        [1900, 2900, 1300, 3260],
        font_size=8.2,
    )
    r.heading("9.2 Controlled smoke-test result", 2)
    r.para(
        "On 20 July 2026, the application was started on 127.0.0.1:8765 with automatic demo traffic and Ollama disabled. After eight seconds, the snapshot contained 140 processed records and 37,125 analysed bytes. The three most recent detector results were a low-severity risky-service alert, a high-severity port-scan alert and a critical SYN-flood alert. The advisor correctly identified the critical item as the first investigation priority."
    )
    r.callout("Verification status", "All eight automated tests passed. Live capture requires local operating-system privileges and should be validated on the presentation machine before the demonstration.", color=BLUE)
    r.heading("9.3 Evaluation against requirements", 2)
    r.para(
        "The MVP meets the defined functional requirements in source and demonstration mode. It provides one local workflow from acquisition to explanation, and it remains usable without an LLM. The primary evaluation gap is production-scale performance and accuracy on a labelled real-world dataset. That work would require authorised data, network-specific baselining and a broader test environment."
    )

    r.heading("10. Project Management and Weekly Work Record", 1, new_page=True)
    r.heading("10.1 Development approach", 2)
    r.para(
        "The internship work followed an iterative approach. Each week produced a concrete result that reduced technical uncertainty: first the scope and data model, then acquisition, detection, presentation, AI integration and verification. Demo mode was treated as a first-class requirement because it allowed progress to be reviewed even when raw-capture privileges or representative traffic were unavailable."
    )
    r.table(
        ["Week", "Focus", "Activities completed", "Deliverable / evidence"],
        [
            ["1", "Problem definition and research", "Reviewed packet-analysis workflows, identified learning needs, defined project boundaries and responsible-use constraints.", "Scope, objectives and initial feature list"],
            ["2", "Architecture and data modelling", "Selected Python/Scapy/FastAPI stack; designed PacketRecord, AlertRecord, store and input-source flow.", "Architecture and module plan"],
            ["3", "Packet acquisition and parsing", "Implemented protocol extraction, address/port handling, direction classification and privacy-conscious metadata.", "Live/PCAP parsing path"],
            ["4", "Backend services", "Built thread-safe store, capture controller, lifecycle API, filters, counters, CSV export and bounded PCAP import.", "Operational local API"],
            ["5", "Threat-detection engine", "Implemented sliding windows, state tracking, thresholds, severity, evidence, cooldowns and recommendations.", "Six explainable detector types"],
            ["6", "Dashboard implementation", "Created responsive interface, status cards, charts, protocol chips, packet drawer and alert list.", "Investigation-oriented web UI"],
            ["7", "AI advisor and resilience", "Integrated Ollama, defensive system prompt, alert context and deterministic bilingual fallback; refined demo traffic.", "Local assistant with no-LLM fallback"],
            ["8", "Testing and documentation", "Added eight automated tests, syntax checks, API smoke test, installation guide, demo guide, architecture notes and report.", "Verified MVP and final documentation"],
        ],
        [700, 1750, 4490, 2420],
        font_size=8.15,
    )
    r.heading("10.2 Milestones", 2)
    r.bullet("M1 — Validated metadata pipeline: synthetic and parsed records enter the same processor.")
    r.bullet("M2 — Explainable alerts: each rule returns evidence and a defensive recommendation.")
    r.bullet("M3 — Usable dashboard: protocol filtering and alert-centred investigation operate from one page.")
    r.bullet("M4 — Resilient AI: application guidance remains available with Ollama disabled.")
    r.bullet("M5 — Reproducible handover: tests, README, architecture and demonstration instructions are included.")

    r.heading("11. Challenges and Solutions", 1, new_page=True)
    r.heading("11.1 Privileged packet capture", 2)
    r.para(
        "Challenge: raw packet capture is intentionally restricted by Linux, and internship review environments may not permit administrator access. Solution: separate capture modes and build a synthetic generator that exercises the exact processing pipeline. The README clearly describes when sudo is required. This preserved both security and demonstrability."
    )
    r.heading("11.2 Alert flooding", 2)
    r.para(
        "Challenge: a continuous scan could satisfy the same threshold on every incoming packet, creating hundreds of duplicate alerts. Solution: use a cooldown dictionary keyed by alert type and source or destination. The first meaningful alert is retained while the interface remains readable."
    )
    r.heading("11.3 Concurrency", 2)
    r.para(
        "Challenge: Scapy callbacks can modify state while the browser reads a snapshot. Solution: guard counters and bounded deques with an RLock, copy current records under the lock and apply slower filtering after release. This produces consistent identifiers and counters without blocking capture for the entire HTTP request."
    )
    r.heading("11.4 AI availability and reliability", 2)
    r.para(
        "Challenge: a local model may be absent, slow or too resource-intensive for the demonstration computer. Solution: isolate the AI call behind SecurityAdvisor, use a timeout and catch failures, then return deterministic guidance. The UI identifies whether Ollama or the fallback produced the answer."
    )
    r.heading("11.5 Balancing evidence and privacy", 2)
    r.para(
        "Challenge: more packet content can improve analysis but also increases privacy and data-handling risk. Solution: restrict the MVP to metadata, make the constraint visible in the dashboard, delete temporary PCAP uploads and bind locally. The trade-off is documented rather than hidden."
    )

    r.heading("12. Skills and Learning Outcomes", 1, new_page=True)
    r.heading("12.1 Technical skills", 2)
    r.bullet("Packet structure and interpretation across Ethernet, IPv4/IPv6, TCP, UDP, DNS, ICMP and ARP.")
    r.bullet("Scapy-based sniffing, PCAP processing and defensive metadata extraction.")
    r.bullet("Sliding-window algorithms, stateful rules, entropy calculation and cooldown design.")
    r.bullet("FastAPI endpoint design, Pydantic validation and application lifecycle management.")
    r.bullet("Thread-safe Python state management with locks, deques and counters.")
    r.bullet("Responsive dashboard development using semantic HTML, CSS and vanilla JavaScript.")
    r.bullet("Local LLM integration, prompt boundaries, timeouts and graceful fallback behaviour.")
    r.bullet("Automated tests, API smoke testing, reproducible installation and technical writing.")
    r.heading("12.2 Professional skills", 2)
    r.para(
        "The project also strengthened requirements analysis, scope control, risk communication and evidence-based reporting. I learned to distinguish a feature that is impressive in a demonstration from one that is reliable enough to include. Building demo mode and a rule fallback required thinking about the evaluator's environment, not only the development environment. Writing limitations and responsible-use guidance improved my ability to communicate uncertainty, which is especially important in cybersecurity."
    )
    r.heading("12.3 Relationship to cybersecurity practice", 2)
    r.para(
        "NetSentry connects several defensive roles. Network analysts use protocol and flow evidence; incident responders prioritise alerts and preserve context; engineers design resilient services; and security governance requires authorised monitoring and controlled data retention. The project did not attempt to master every role, but it demonstrated how their concerns interact in one system."
    )

    r.heading("13. Ethics, Limitations and Future Work", 1, new_page=True)
    r.heading("13.1 Responsible use", 2)
    r.para(
        "Packet monitoring must be performed only on networks and devices that the operator owns or has explicit permission to inspect. Even without payloads, IP addresses, ports, DNS names and timing can reveal sensitive relationships. Exports should be access-controlled and deleted when no longer required. The application intentionally defaults to localhost and presents privacy messaging beside the capture controls."
    )
    r.heading("13.2 Current limitations", 2)
    r.bullet("Fixed thresholds may be too sensitive for busy networks and too permissive for quiet ones.")
    r.bullet("Encrypted application payloads are not inspected, and protocol identification is mainly port-based.")
    r.bullet("The in-memory store is not designed for long-term retention or distributed sensors.")
    r.bullet("The interface polls once per second and is not benchmarked for enterprise packet rates.")
    r.bullet("No authentication is implemented because the service is intended for localhost only.")
    r.bullet("The advisor can produce incomplete or incorrect suggestions and always requires human verification.")
    r.heading("13.3 Future development roadmap", 2)
    r.table(
        ["Priority", "Improvement", "Expected value"],
        [
            ["Near term", "Configuration file for thresholds, allowlists and retention", "Supports different lab and network profiles."],
            ["Near term", "PCAP export and session tagging", "Improves reproducible investigation and evidence grouping."],
            ["Medium", "Flow aggregation and TCP-state tracking", "Reduces record volume and improves scan/flood accuracy."],
            ["Medium", "Baseline learning by host and hour", "Replaces universal thresholds with environment-aware detection."],
            ["Medium", "Authentication and TLS for non-local deployment", "Protects dashboards exposed beyond localhost."],
            ["Long term", "Zeek/Suricata integration and SIEM export", "Uses established sensors and central correlation."],
            ["Long term", "Labelled-dataset evaluation", "Measures precision, recall, false-positive rate and performance."],
        ],
        [1250, 3750, 4360],
        font_size=8.6,
    )

    r.heading("14. Conclusion", 1, new_page=True)
    r.para(
        "The internship project achieved its main goal: to build a working network-traffic monitoring application that combines protocol exploration, suspicious-activity detection and practical defensive guidance. NetSentry AI can capture live metadata, analyse PCAP files, generate safe synthetic traffic, filter protocols, explain six categories of suspicious behaviour and export results. Its local AI assistant adds context without becoming a dependency for detection or operation."
    )
    r.para(
        "The most important design outcome is the separation of responsibilities. Inputs share one processor; metadata is separate from payloads; detection is separate from language generation; and the browser receives a validated local API. This structure made the system easier to test and allowed graceful behaviour when privileges or Ollama were unavailable."
    )
    r.para(
        "The controlled evaluation demonstrated that the application starts locally, processes synthetic traffic, returns protocol statistics, creates expected findings and gives a prioritised rule-based answer. All eight automated tests passed. The project remains an educational MVP, but it provides a solid base for future work on baselines, flow aggregation, established IDS integration and measurable detection accuracy."
    )
    r.callout("Final reflection", "The project developed practical confidence in turning raw network observations into explainable, testable and responsibly presented security information.")

    r.heading("References", 1, new_page=True)
    references = [
        "Scapy Project. Scapy 2.7.1 Documentation: Introduction and Usage. https://scapy.readthedocs.io/en/latest/ (accessed 20 July 2026).",
        "FastAPI. Official Documentation and Tutorial. https://fastapi.tiangolo.com/ (accessed 20 July 2026).",
        "Ollama. API Introduction. https://docs.ollama.com/api/introduction (accessed 20 July 2026).",
        "National Institute of Standards and Technology. Guide to Intrusion Detection and Prevention Systems (IDPS), NIST SP 800-94. https://csrc.nist.gov/pubs/sp/800/94/final",
        "Wireshark Foundation. Wireshark User's Guide. https://www.wireshark.org/docs/wsug_html_chunked/",
        "Python Software Foundation. Python 3 Documentation: socket, threading, collections and unittest. https://docs.python.org/3/",
        "OWASP Foundation. Logging Cheat Sheet. https://cheatsheetseries.owasp.org/cheatsheets/Logging_Cheat_Sheet.html",
    ]
    for index, item in enumerate(references):
        r.numbered(item, restart=index == 0)

    r.heading("Appendix A — Local Installation and Demonstration", 1, new_page=True)
    r.heading("A.1 Install on Kali Linux or Ubuntu", 2)
    r.code("sudo apt update\nsudo apt install -y python3-venv libpcap-dev\npython3 -m venv .venv\nsource .venv/bin/activate\npip install -e .")
    r.heading("A.2 Recommended demonstration", 2)
    r.code("netsentry --demo")
    r.para("Open http://127.0.0.1:8000 in a browser. The demo requires no root privileges. Allow approximately ten seconds for the first port-scan and SYN-flood indicators to appear.")
    r.numbered("Show total packets, rate, bytes and engine status.", restart=True)
    r.numbered("Select DNS, TLS and TCP protocol chips to demonstrate filtering.")
    r.numbered("Search for 10.10.5.44 to isolate the synthetic scanning source.")
    r.numbered("Enable Alerts only and open a highlighted packet.")
    r.numbered("Choose Ask AI on a high or critical alert and submit the prepared question.")
    r.numbered("Explain the rules fallback if Ollama is not installed.")
    r.numbered("Stop capture and export the metadata as CSV.")
    r.heading("A.3 Live capture", 2)
    r.code("sudo .venv/bin/netsentry")
    r.para("In the interface, choose Live interface, select eth0, wlan0 or another authorised adapter, and click Start capture. Generate permitted traffic by opening a website or using ping in a separate terminal. Never monitor a network without authorisation.")
    r.heading("A.4 Optional local model", 2)
    r.code("ollama pull qwen2.5:3b\nNETSENTRY_USE_OLLAMA=1 NETSENTRY_OLLAMA_MODEL=qwen2.5:3b netsentry --demo")

    r.heading("Appendix B — Project File Structure", 1, new_page=True)
    r.code(
        "netsentry/\n"
        "  main.py            FastAPI routes and lifecycle\n"
        "  capture.py         Live, demo and PCAP inputs\n"
        "  packet_parser.py   Scapy metadata conversion\n"
        "  detector.py        Explainable detection rules\n"
        "  advisor.py         Ollama and offline guidance\n"
        "  store.py           Thread-safe bounded state\n"
        "  static/            Dashboard CSS and JavaScript\n"
        "  templates/         Dashboard HTML\n"
        "tests/               Detector and pipeline tests\n"
        "docs/                Architecture and demo guide"
    )
    r.heading("B.1 Main configuration variables", 2)
    r.table(
        ["Variable", "Default", "Purpose"],
        [
            ["NETSENTRY_AUTO_DEMO", "0", "Starts synthetic capture during application startup when set to 1."],
            ["NETSENTRY_USE_OLLAMA", "1", "Set to 0 to force deterministic advisor responses."],
            ["NETSENTRY_OLLAMA_URL", "http://127.0.0.1:11434", "Local or explicitly configured Ollama API base URL."],
            ["NETSENTRY_OLLAMA_MODEL", "qwen2.5:3b", "Model requested from Ollama."],
        ],
        [2700, 2100, 4560],
        font_size=8.6,
    )
    r.heading("B.2 Re-run verification", 2)
    r.code("python -m unittest discover -s tests -v\npython -m compileall -q netsentry tests\nnode --check netsentry/static/app.js")

    r.doc.save(REPORT_PATH)
    return REPORT_PATH


if __name__ == "__main__":
    print(build())
